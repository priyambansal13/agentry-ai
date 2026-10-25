"""Invoices: import (CSV/Excel, any column names), payments and matching, Razorpay payment links.

Import maps columns heuristically (covers Google Sheets, Tally "Bills Receivable" and QuickBooks "Open Invoices"
exports) and asks the fast model only when required columns can't be identified. Every import is previewed before
it is committed.
"""
import csv
import io
import json
import re
from datetime import UTC, date, datetime
from decimal import Decimal, InvalidOperation

from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncConnection

from agentry.connectors.http import client
from agentry.connectors.store import get_connection
from agentry.events import emit
from agentry.leads import resolve_contact
from agentry.llm.gateway import AllModelsFailed, complete

FIELDS = {
    "number": ["invoice", "invoice no", "invoice number", "inv", "num", "ref. no.", "ref no", "reference", "bill no", "voucher no"],
    "company": ["company", "customer", "party's name", "party name", "client", "customer name", "account"],
    "contact_name": ["contact", "contact person", "name", "attention"],
    "email": ["email", "customer email", "e-mail", "email address"],
    "phone": ["phone", "mobile", "whatsapp", "contact number"],
    "amount": ["amount", "opening amount", "total", "invoice amount", "bill amount", "gross"],
    "balance": ["open balance", "pending amount", "balance", "balance due", "outstanding", "amount due"],
    "amount_paid": ["amount paid", "paid", "received"],
    "issued_on": ["invoice date", "date", "issued", "bill date"],
    "due_on": ["due date", "due on", "due", "payment due"],
    "status": ["status", "state"],
    "currency": ["currency", "cur"],
    "whatsapp_consent": ["whatsapp consent", "whatsapp opt-in", "whatsapp optin", "whatsapp ok", "whatsapp reminders"],
}
REQUIRED = ("number", "amount", "due_on")


def parse_file(data: bytes, filename: str) -> list[dict]:
    if filename.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
        rows = list(ws.iter_rows(values_only=True))
        headers = [str(h).strip() if h is not None else "" for h in rows[0]]
        return [{headers[i]: v for i, v in enumerate(r) if i < len(headers) and headers[i]} for r in rows[1:] if any(r)]
    textdata = data.decode("utf-8-sig", errors="replace")
    return [dict(r) for r in csv.DictReader(io.StringIO(textdata)) if any((v or "").strip() for v in r.values())]


def _norm(h: str) -> str:
    return re.sub(r"\s+", " ", h.strip().lower())


def map_columns_heuristic(headers: list[str]) -> dict[str, str]:
    mapping: dict[str, str] = {}
    used: set[str] = set()
    for field, names in FIELDS.items():  # exact synonym first, in priority order
        for name in names:
            h = next((h for h in headers if _norm(h) == name and h not in used), None)
            if h:
                mapping[field], _ = h, used.add(h)
                break
    return mapping


async def map_columns(org_id: str, headers: list[str], sample: list[dict]) -> tuple[dict[str, str], str]:
    m = map_columns_heuristic(headers)
    if all(f in m for f in REQUIRED):
        return m, "heuristic"
    prompt = [{"role": "system", "content": "Map spreadsheet columns to invoice fields. Return only JSON {field: column} "
               f"using fields {list(FIELDS)} and columns exactly as given; omit fields with no matching column."},
              {"role": "user", "content": json.dumps({"columns": headers, "sample_rows": sample[:3]}, default=str)}]
    try:
        res = await complete("fast", prompt, org_id=org_id, prompt_key="invoice_column_mapper", prompt_version="v1",
                             json_output=True, mock_response=json.dumps(m))
        ai = {k: v for k, v in (res.data or {}).items() if k in FIELDS and v in headers}
        return {**ai, **m}, "ai" if not res.mock else "heuristic"
    except AllModelsFailed:
        return m, "heuristic"


def _num(v) -> Decimal | None:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float, Decimal)):
        return Decimal(str(v))
    s = re.sub(r"[^\d.\-]", "", str(v).replace(",", ""))
    try:
        return Decimal(s) if s not in ("", "-", ".") else None
    except InvalidOperation:
        return None


def _date(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v or "").strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y", "%d %b %Y", "%d-%b-%Y", "%d %B %Y", "%b %d, %Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def normalise_rows(rows: list[dict], mapping: dict[str, str], default_currency: str) -> tuple[list[dict], list[str]]:
    out, issues = [], []
    for i, r in enumerate(rows, start=2):  # row 1 is the header
        get = lambda f, r=r: r.get(mapping[f]) if f in mapping else None
        number, amount, due = str(get("number") or "").strip(), _num(get("amount")), _date(get("due_on"))
        if not number or amount is None or due is None:
            issues.append(f"row {i}: missing " + ", ".join(f for f, v in (("invoice number", number), ("amount", amount), ("due date", due)) if not v))
            continue
        balance, paid = _num(get("balance")), _num(get("amount_paid"))
        if paid is None:
            paid = amount - balance if balance is not None else Decimal(0)
        status_raw = str(get("status") or "").lower()
        status = "paid" if (paid >= amount or "paid" == status_raw) else "partially_paid" if paid > 0 else "open"
        out.append({"number": number, "amount": amount, "amount_paid": min(paid, amount), "due_on": due,
                    "issued_on": _date(get("issued_on")), "status": status,
                    "currency": str(get("currency") or default_currency).upper()[:3],
                    "company": (str(get("company")).strip() if get("company") else None),
                    "contact_name": (str(get("contact_name")).strip() if get("contact_name") else None),
                    "email": (str(get("email")).strip() if get("email") else None),
                    "phone": (str(get("phone")).strip() if get("phone") else None),
                    # #16: WhatsApp reminders need the customer's agreement (e.g. in your terms); a yes in this column records it
                    "whatsapp_consent": str(get("whatsapp_consent") or "").strip().lower() in ("yes", "y", "true", "1", "agreed")})
    return out, issues


async def upsert_invoices(conn: AsyncConnection, org_id: str, rows: list[dict], source: str) -> dict:
    created = updated = 0
    ids = []
    for r in rows:
        contact_id = None
        if r["email"] or r["phone"]:
            contact_id = await resolve_contact(conn, org_id, name=r["contact_name"], email=r["email"], phone=r["phone"],
                                               company=r["company"], source=f"import:{source}")
            if r.get("whatsapp_consent") and r["phone"]:
                from agentry.leads import record_consent
                await record_consent(conn, org_id, contact_id, "whatsapp", True, source=f"import:{source}")
        elif r["company"]:  # e.g. Tally exports: the party's contact details live in the ledger master
            contact_id = (await conn.execute(text("""select id from app.contacts where lower(company) = lower(:c)
                and merged_into is null order by created_at limit 1"""), {"c": r["company"]})).scalar_one_or_none()
        row = (await conn.execute(text("""
            insert into app.invoices (org_id, contact_id, number, currency, amount, amount_paid, issued_on, due_on, status, source)
            values (:o, :k, :n, :cur, :a, :p, :i, :d, :s, :src)
            on conflict (org_id, number) do update set amount = excluded.amount,
              amount_paid = greatest(app.invoices.amount_paid, excluded.amount_paid), due_on = excluded.due_on,
              contact_id = coalesce(app.invoices.contact_id, excluded.contact_id),
              status = case when app.invoices.status in ('disputed','written_off','void') then app.invoices.status
                            else excluded.status end, updated_at = now()
            returning id, (xmax = 0) as inserted"""),
            {"o": org_id, "k": contact_id, "n": r["number"], "cur": r["currency"], "a": r["amount"], "p": r["amount_paid"],
             "i": r["issued_on"], "d": r["due_on"], "s": r["status"], "src": source})).first()
        ids.append(str(row.id))
        created, updated = created + bool(row.inserted), updated + (not row.inserted)
    return {"created": created, "updated": updated, "invoice_ids": ids}


async def enroll_open(conn: AsyncConnection, org_id: str) -> int:
    """Put every open, contactable invoice into the reminder sequence (idempotent)."""
    from agentry.sequences import enroll
    status = (await conn.execute(text("select status from app.automation_instances where template_key = 'invoice_follow_up'"))).scalar_one_or_none()
    if status not in ("live", "approval_mode"):
        return 0
    rows = (await conn.execute(text("""select id, contact_id from app.invoices where status in ('open','partially_paid')
        and contact_id is not null"""))).all()
    n = 0
    for r in rows:
        n += bool(await enroll(conn, org_id, sequence_key="invoice_default", subject_type="invoice",
                               subject_id=str(r.id), contact_id=str(r.contact_id)))
    return n


# ---------- payments ----------

async def apply_payment(conn: AsyncConnection, org_id: str, *, invoice_id: str | None, amount: Decimal, currency: str,
                        gateway: str, reference: str, match: str, raw: dict | None = None,
                        received_at: datetime | None = None) -> str | None:
    """Record a payment once (gateway + reference is unique) and update the invoice. Returns the payment id, or
    None if this payment was already recorded."""
    pid = (await conn.execute(text("""insert into app.payments (org_id, invoice_id, amount, currency, gateway, reference, match, raw, received_at)
        values (:o, :i, :a, :c, :g, :r, :m, cast(:raw as jsonb), coalesce(:at, now()))
        on conflict (org_id, gateway, reference) do nothing returning id"""),
        {"o": org_id, "i": invoice_id, "a": amount, "c": currency, "g": gateway, "r": reference, "m": match,
         "raw": json.dumps(raw or {}, default=str), "at": received_at})).scalar_one_or_none()
    if pid is None:
        return None
    if invoice_id:
        current = (await conn.execute(text("select number, status from app.invoices where id = cast(:i as uuid)"), {"i": invoice_id})).first()
        if current and current.status in ("paid", "void", "written_off"):  # #17: never fold a payment silently into a closed invoice
            await conn.execute(text("""insert into app.approvals (org_id, kind, automation_key, title, body, payload)
                values (:o, 'payment_review', 'invoice_follow_up', :t, :b, cast(:p as jsonb))"""),
                {"o": org_id, "t": f"Payment of {currency} {amount:,.2f} for invoice {current.number}, which is already {current.status}",
                 "b": "Check whether this is a duplicate, an advance, or meant for another invoice.",
                 "p": json.dumps({"payment_id": str(pid), "reference": reference, "invoice_id": invoice_id})})
            await conn.execute(text("update app.payments set invoice_id = null where id = :p"), {"p": pid})
            return str(pid)
        inv = (await conn.execute(text("""update app.invoices set amount_paid = least(amount, amount_paid + :a),
            status = case when amount_paid + :a >= amount then 'paid' else 'partially_paid' end, updated_at = now()
            where id = :i returning number, status, amount - amount_paid as balance"""), {"a": amount, "i": invoice_id})).first()
        await withdraw_reminder_cards(conn, invoice_id, "paid" if inv.status == "paid" else "part paid: the next reminder shows the new balance")
        await emit(conn, org_id, "payment.received", {"amount": float(amount), "currency": currency, "invoice_id": invoice_id,
                                                      "payment_id": str(pid), "reference": reference,
                                                      "invoice": inv.number, "status": inv.status, "match": match,
                                                      "gateway": gateway})
    else:
        await conn.execute(text("""insert into app.approvals (org_id, kind, automation_key, title, body, payload)
            values (:o, 'payment_review', 'invoice_follow_up', :t, :b, cast(:p as jsonb))"""),
            {"o": org_id, "t": f"Unmatched payment {currency} {amount:,.2f} ({gateway})",
             "b": "We couldn't tell which invoice this payment is for. Pick the invoice or mark it as not ours.",
             "p": json.dumps({"payment_id": str(pid), "reference": reference})})
    return str(pid)


async def match_invoice(conn: AsyncConnection, *, invoice_hint: str | None, amount: Decimal, email: str | None,
                        phone: str | None) -> tuple[str | None, str]:
    """Strong: the payment names the invoice. Weak: one open invoice of that customer with exactly that balance."""
    if invoice_hint:
        row = (await conn.execute(text("select id from app.invoices where id::text = :h or number = :h"), {"h": invoice_hint})).first()
        if row:
            return str(row.id), "strong"
    from agentry.leads import normalize_email, normalize_phone
    em, ph = normalize_email(email), normalize_phone(phone)
    rows = (await conn.execute(text("""select i.id from app.invoices i join app.contacts k on k.id = i.contact_id
        where i.status in ('open','partially_paid') and i.amount - i.amount_paid = :a
          and ((cast(:e as text) is not null and k.email_norm = :e) or (cast(:p as text) is not null and k.phone_e164 = :p))"""),
        {"a": amount, "e": em, "p": ph})).all()
    if len(rows) == 1:
        return str(rows[0].id), "weak"
    return None, "unmatched"


async def payment_link(conn: AsyncConnection, org_id: str, invoice: dict) -> str | None:
    """A payment link for the invoice's current balance (#4, #15, #19): Razorpay for INR, Stripe for other currencies
    (needs the client's Stripe secret key). Made once, and again only when the balance changed (a part payment), so a
    customer is never asked for more than they owe. The link carries the invoice id, so the payment matches exactly."""
    balance = Decimal(str(invoice["amount"])) - Decimal(str(invoice["amount_paid"]))
    if balance <= 0:
        return None
    if invoice.get("payment_link") and invoice.get("payment_link_amount") is not None and Decimal(str(invoice["payment_link_amount"])) == balance:
        return invoice["payment_link"]
    notes = {"invoice_id": str(invoice["id"]), "org_id": org_id, "invoice": invoice["number"]}
    link, provider = None, None
    rz = await get_connection(org_id, "razorpay") if invoice["currency"] == "INR" else None
    st = await get_connection(org_id, "stripe") if not rz else None
    async with client() as c:
        if rz:
            r = await c.post("https://api.razorpay.com/v1/payment_links", auth=(rz["secret"]["key_id"], rz["secret"]["key_secret"]),
                             json={"amount": int(balance * 100), "currency": invoice["currency"],
                                   "description": f"Invoice {invoice['number']}", "reference_id": f"{invoice['number'][:30]}-{int(balance)}",
                                   "notes": notes,
                                   "customer": {k: v for k, v in {"name": invoice.get("contact_name"), "email": invoice.get("email"),
                                                                  "contact": invoice.get("phone_e164")}.items() if v}})
            link, provider = (r.json().get("short_url"), "razorpay") if r.status_code < 300 else (None, None)
        elif st and st["secret"].get("secret_key"):
            key = (st["secret"]["secret_key"], "")
            price = await c.post("https://api.stripe.com/v1/prices", auth=key, data={
                "currency": invoice["currency"].lower(), "unit_amount": int(balance * 100), "product_data[name]": f"Invoice {invoice['number']}"})
            if price.status_code < 300:
                r = await c.post("https://api.stripe.com/v1/payment_links", auth=key, data={
                    "line_items[0][price]": price.json()["id"], "line_items[0][quantity]": 1,
                    **{f"metadata[{k}]": v for k, v in notes.items()}, **{f"payment_intent_data[metadata][{k}]": v for k, v in notes.items()}})
                link, provider = (r.json().get("url"), "stripe") if r.status_code < 300 else (None, None)
    if link:
        await conn.execute(text("""update app.invoices set payment_link = :l, payment_link_amount = :a, payment_link_provider = :p
            where id = cast(:i as uuid)"""), {"l": link, "a": balance, "p": provider, "i": str(invoice["id"])})
    return link


async def withdraw_reminder_cards(conn: AsyncConnection, invoice_id: str, why: str) -> int:
    """Reminder drafts still waiting in Needs you are withdrawn when the invoice is paid (or part paid, disputed, void):
    approving one would ask the customer for money they no longer owe (#5)."""
    r = await conn.execute(text("""update app.approvals a set status = 'resolved', decided_at = now(),
          payload = a.payload || jsonb_build_object('withdrawn', cast(:w as text))
        from app.sequence_enrollments e
        where a.kind = 'outbound_draft' and a.status = 'pending' and e.id::text = a.payload->>'enrollment_id'
          and e.subject_type = 'invoice' and e.subject_id = cast(:i as uuid)"""), {"i": invoice_id, "w": why})
    if not why.startswith("part paid"):  # "time for a call" is moot once it's paid, disputed or closed
        await conn.execute(text("""update app.approvals set status = 'resolved', decided_at = now(),
              payload = payload || jsonb_build_object('withdrawn', cast(:w as text))
            where kind = 'escalation' and status = 'pending' and payload->>'invoice_id' = :i and title like 'Reminders finished%'"""),
            {"i": invoice_id, "w": why})
    return r.rowcount


def now_utc() -> datetime:
    return datetime.now(UTC)
