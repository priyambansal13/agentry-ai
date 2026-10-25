"""Knowledge from pasted text and uploaded files (PDF, Word, text, CSV), searchable by the agents."""
import io

from conftest import expect
from docx import Document

from agentry.knowledge import search

ALL = ("owner", "manager", "staff", "finance", "viewer", "agentry")
ADMINS = {"owner": 201, "manager": 201, "staff": 403, "finance": 403, "viewer": 403, "agentry": 201}


def tiny_pdf(text: str) -> bytes:
    """A minimal one-page PDF with a line of text (valid enough for pypdf)."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
            b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out, offsets = io.BytesIO(), []
    out.write(b"%PDF-1.4\n")
    for i, o in enumerate(objs, 1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n".encode() + o + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode())
    return out.getvalue()


def word_doc() -> bytes:
    d = Document()
    d.add_paragraph("Onboarding takes five working days after the kickoff call.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Plan", "Monthly price"
    t.cell(1, 0).text, t.cell(1, 1).text = "Growth plan", "Rs 24,999"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


async def test_paste_text_and_upload_files(client, team_org):
    oid, h = await team_org(["inbox_triage"])
    got = {r: (await client.post("/org/knowledge/text", headers=h[r], json={
        "title": f"Support hours ({r})", "content": "Our support desk is open Monday to Saturday, 9am to 7pm IST."})).status_code for r in ALL}
    expect(got, **ADMINS)

    files = {
        "pricing.pdf": ("application/pdf", tiny_pdf("Starter plan costs Rs 9,999 per month")),
        "plans.docx": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", word_doc()),
        "refunds.txt": ("text/plain", b"Refund policy: full refund within 14 days of purchase, no questions asked."),
        "integrations.csv": ("text/csv", b"tool,supported\nSlack,yes\nHubSpot,yes\nSalesforce,coming soon\n"),
    }
    for name, (mime, data) in files.items():
        r = await client.post("/org/knowledge/file", headers=h["owner"], files={"file": (name, data, mime)})
        assert r.status_code == 201, (name, r.text)
    docs = {d["title"]: d for d in (await client.get("/org/knowledge", headers=h["owner"])).json()["items"]}
    assert set(files) <= set(docs)

    # the agents find it (offline embeddings in tests: keyword-level search)
    assert any("9,999" in x.content for x in await search(oid, "Starter plan costs per month"))
    assert any("Rs 24,999" in x.content for x in await search(oid, "Growth plan monthly price"))
    assert any("14 days" in x.content for x in await search(oid, "refund policy"))

    # re-uploading the same file name replaces it
    await client.post("/org/knowledge/file", headers=h["owner"],
                      files={"file": ("refunds.txt", b"Refund policy: full refund within 30 days of purchase.", "text/plain")})
    refunds = [d for d in (await client.get("/org/knowledge", headers=h["owner"])).json()["items"] if d["title"] == "refunds.txt"]
    assert len(refunds) == 1

    # clear errors for bad uploads; front desk can't upload
    bad = await client.post("/org/knowledge/file", headers=h["owner"], files={"file": ("photo.png", b"\x89PNG....", "image/png")})
    assert bad.status_code == 422 and "PDF" in bad.json()["detail"]
    empty = await client.post("/org/knowledge/file", headers=h["owner"], files={"file": ("blank.txt", b"   ", "text/plain")})
    assert empty.status_code == 422
    broken = await client.post("/org/knowledge/file", headers=h["owner"], files={"file": ("broken.pdf", b"%PDF-1.4 garbage", "application/pdf")})
    assert broken.status_code == 422
    assert (await client.post("/org/knowledge/file", headers=h["staff"], files={"file": ("x.txt", b"hello world, some text", "text/plain")})).status_code == 403
